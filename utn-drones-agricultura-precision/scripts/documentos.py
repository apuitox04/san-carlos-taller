"""Generación reproducible de documentos originales. No contiene datasets ajenos."""
from contenido import *
import json, csv, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY='264C6E'; TEAL='00918C'
def base(titulo, subtitulo):
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.8)
    for st in d.styles:
        for b in list(st.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
    for name,size in [('Normal',10.5),('Title',24),('Heading 1',17),('Heading 2',12)]:
        s=d.styles[name]; s.font.name='Fira Sans'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string('264C6E' if name!='Normal' else '263B48')
        s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.1
    foot=sec.footer.paragraphs[0]; foot.text='UTN · Administración Agroindustrial | Material de trabajo · 18 octubre 2026'
    foot.style='Caption'; foot.runs[0].font.size=Pt(8)
    d.add_heading(titulo,0); d.add_paragraph(subtitulo)
    return d
def table(d,headers,rows):
    t=d.add_table(rows=1, cols=len(headers)); t.style='Light Shading Accent 1'
    for c,txt in zip(t.rows[0].cells,headers):
        c.text=str(txt); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),NAVY); c._tc.get_or_add_tcPr().append(shade)
        for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=str(txt)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(9)
    d.add_paragraph('')
    return t
def save(d,path):
    p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);d.save(p)
def guide(title,subtitle,pages,path,refs):
    d=base(title,subtitle)
    for i,(h,intro,steps) in enumerate(pages):
        if i:d.add_page_break()
        d.add_heading(f'{i+1}. {h}',1);d.add_paragraph(intro)
        for label,text in steps:
            d.add_heading(label,2);d.add_paragraph(text)
        d.add_paragraph('Fuentes: '+refs,style='Caption')
    save(d,path)

d=base('Programa del taller','Drones y agricultura de precisión · 120 minutos · Modalidad demostrativa')
d.add_paragraph('Población: estudiantes de Administración Agroindustrial de la Universidad Técnica Nacional, Costa Rica. Objetivo: relacionar productos geoespaciales, controles y decisiones de gestión. El taller no forma operadores especializados ni autoriza prescripciones.')
table(d,['Minutos','Actividad','Diapositivas'],[(x[0],x[2],x[1]) for x in SECCIONES])
d.add_heading('Condiciones para impartirlo',1)
d.add_paragraph('Usar proyectos previamente procesados y ensayados. Si faltan, impartir la alternativa documental identificada en el guion. No iniciar procesamiento completo en aula. Se conserva la distribución solicitada por la UTN: 15 + 15 + 30 + 30 + 15 + 10 + 5 = 120 minutos.')
d.add_paragraph('Fecha, hora, lugar, inscripción y expositor: por confirmar. Programa de trabajo; no constituye aprobación institucional.')
save(d,'taller/programa/Programa_120_minutos.docx')

d=base('Guion del instructor','Instrumento operativo · Correspondencia con las 21 diapositivas')
for i,sec in enumerate(SECCIONES):
    if i:d.add_page_break()
    tiempo,slides,titulo,obj,exp,prog,ops,res,preg,concept,alt=sec
    d.add_heading(f'{tiempo} min | {titulo}',1)
    d.add_paragraph('Diapositivas '+slides+' · Mantener el reloj visible.')
    for h,txt in [('Objetivo',obj),('Explicación breve',exp),('Programa',prog),('Operaciones',ops),('Resultado que se muestra',res),('Pregunta a estudiantes',preg),('Concepto clave',concept),('Alternativa técnica',alt)]:
        d.add_heading(h,2);d.add_paragraph(txt)
save(d,'taller/guion-instructor/Guion_instructor.docx')
guide('Guía D · PIX4Dmatic Pro','Ejercicio documental · Procesamiento y ensayo pendientes',MATIC,'ejercicios/pix4dmatic-pro/Guia_PIX4Dmatic.docx','R02, R03, R05, R07, M02. URLs en REFERENCIAS.md.')
guide('Guía E1 · Índices de vegetación','Field trial · Ejemplo oficial de capacitación',INDICES,'ejercicios/pix4dfields/Guia_indices.docx','M01 pp. 318–333; R06, R09, R10, R13. URLs en REFERENCIAS.md.')
guide('Guía E2 · Malezas y Magic Tool','Targeted Operations · Ejemplo oficial de capacitación',MALEZAS,'ejercicios/pix4dfields/Guia_malezas.docx','M01 pp. 300–317; R11, R12, R16, R17. URLs en REFERENCIAS.md.')
guide('Propuesta de curso formal','Drones y tecnologías geoespaciales aplicadas a la agroindustria · 40 horas · Pendiente de validación UTN',CURSO,'propuesta-curso/docx/Propuesta_curso_40h.docx','R01, R02, R06, R13, R15, R19 y manuales aportados. Presupuesto: simulación explícita.')
d=base('Ficha de decisión y salida','Actividad de Administración Agroindustrial · 10 minutos y cierre de 2 minutos')
for h,txt in [('Problema','¿Qué decisión necesita tomar en una finca o empresa?'),('Producto','¿RGB, nube, DSM, NDVI, NDRE o selección de malezas? Justifique.'),('Comprobación','¿Qué observación independiente hace falta?'),('Gestión','¿Quién comprueba, cuándo y con qué costo?'),('Beneficio','¿Qué indicador medirá antes y después? No inventar ahorros.'),('Salida','Escriba un producto, una limitación y la comprobación siguiente.')]:
    d.add_heading(h,2); d.add_paragraph(txt); d.add_paragraph('____________________________________________________________________')
save(d,'taller/evaluacion/Ficha_decision.docx')

revision='''# Investigación y selección de contenidos

Consulta documental: 17 octubre 2026. Reconstrucción de archivos: 18 octubre 2026.

## Pertinencia académica
El perfil UTN vincula gestión agroindustrial, tecnología y optimización de recursos (R01). Se seleccionan productos para decisiones y controles, no una formación de pilotaje. MAG documenta adopción y extensión en agricultura de precisión (R19–R21). Café, piña, caña, pasturas e inventarios son contextos posibles; su uso local requiere datos y acuerdos. No se atribuyen ahorros.

## Manuales aportados
M01 Fields Essentials, 487 páginas: fundamentos, límites, índices, radiometría, estadísticas y dos casos. Targeted Operations ocupa páginas PDF 300–317 y Field trial 318–333. M02 Matic Essentials Standard 2026, 313 páginas: entradas, coordenadas, control, calibración y productos. M03 mapper Essentials y M04 mapper In-depth: fundamentos y antecedentes; sus botones y menús no se trasladan a Matic. No se redistribuyen los cuadernos.

## Plataformas y versiones
Matic Standard y Pro reconstruyen fotografías: calibración, nube, malla, DSM y ortomosaico; Pro añade funciones topográficas (R02). Referencia estable 2.10.3; 2.12.0 se identifica como preview (R03). Fields 2.13.5: procesamiento agrícola, índices, límites, estadísticas y operaciones (R04, R06). No se confunde procesamiento rápido con comprobación geométrica independiente. Fields no admite GCP y su licencia de prueba bloquea exportación (R06, R15). Las versiones son referencias documentales, no instalaciones ensayadas.

## Interoperabilidad
Matic exporta GeoTIFF de ortomosaico y DSM, LAS/LAZ/XYZ, OBJ e informe PDF (R05). Fields admite imágenes y GeoTIFF y produce índices y salidas operativas según licencia y equipo (R06). Un ortomosaico RGB no crea bandas Red Edge o NIR. DATA.TIF conserva valores; el raster coloreado sirve para presentación. El destino debe conservar CRS, unidades, bandas, NoData y fecha. La apertura en un receptor no certifica compatibilidad de maquinaria.

## Captura y conceptos
Mavic 3M: RGB 20 MP y cuatro cámaras 5 MP G R RE NIR; se adopta ficha técnica R14 ante discrepancias de la FAQ. Separar RGB y multiespectral, desactivar Dewarping en RGB y revisar nadir y radiometría (R13). GSD no equivale a exactitud. Las alturas elipsoidales y ortométricas no se intercambian sin transformación documentada.

NDVI=(NIR−Red)/(NIR+Red); NDRE=(NIR−Red Edge)/(NIR+Red Edge). Son adimensionales. Con reflectancias no negativas y denominador distinto de cero, el rango teórico es [−1,1]. Enmascarar NoData y denominador cero. No existen umbrales universales que sustituyan diagnóstico de campo. No comparar fechas sin controles radiométricos y fenológicos.

## Casos y certeza de la evidencia
1. Ejemplo oficial de capacitación Matic100: Suiza, 100 RGB AeriaX, eBee X, tres GCP, EPSG 32632 y altura elipsoidal WGS84. No es un levantamiento agrícola UTN.
2. Ejemplo oficial Field trial: Alemania, trigo, Bluegrass/Sequoia, 1328 imágenes, 50 m, 5.1 cm GSD, aproximadamente 1 ha (M01). Rostock y Winter wheat públicos son otros conjuntos; la ficha Winter wheat presenta una discrepancia de cámara. No se mezclan sus resultados.
3. Ejemplo oficial Targeted Operations: zarzamora en pastura australiana, XAG M500 RGB, 254 imágenes, 60 m, 3.15 cm GSD y aproximadamente 10 ha (M01). Rumex con Mavic 3M es otro caso (R18).
4. Simulaciones didácticas: diagrama de parcelas, cuadrícula Magic, fórmulas interactivas y presupuesto. No son resultados calculados de los casos.
5. Propuestas pendientes: prácticas con productores de Costa Rica, curso, presupuesto final y aprobación UTN.

## Selección y viabilidad en 120 minutos
Conservar el horario solicitado. Matic 30 min, índices 30 min y Magic 15 min solo son viables con proyectos listos. El guion limita exposición conceptual y reserva tiempo para decisiones. Descargar y procesar antes. No existe ensayo cronometrado en esta producción: registrar versión, licencia, datos y duración en las plantillas antes de anunciar demostraciones operativas. Sin ello usar la alternativa documental, identificándola.

## Límites de esta producción
No se ejecutaron aplicaciones Pix4D ni proyectos nativos. Las vistas incluidas son imágenes públicas oficiales, no capturas tomadas de una instalación vigente. No se fabrican informes, índices, estadísticas ni exportaciones. Las guías documentan rutas del material de capacitación que deben confirmarse en la versión instalada. El requisito de proyecto previamente procesado sigue pendiente.
'''
(ROOT/'investigacion/revision-tecnica.md').write_text(revision)
refs='# Referencias verificadas\n\nConsulta: 17 octubre 2026. Los manuales aportados no se redistribuyen.\n\n'
for code,inst,title,url,claim in REFS:refs+=f'## {code} · {title}\n\nInstitución o autor: {inst}.\n\nURL: {url}\n\nConsulta: {FECHA}. Afirmaciones respaldadas: {claim}\n\n'
(ROOT/'REFERENCIAS.md').write_text(refs);(ROOT/'investigacion/fuentes-verificadas.md').write_text(refs)
assets='''# Procedencia y derechos de recursos

No se incluyen fotografías de participantes ni logotipos UTN. La mención institucional identifica el público; no acredita aprobación.

| Archivo | Procedencia | Uso y crédito |
|---|---|---|
| fotografias/mavic3m.jpg | R22, ZLEA, Wikimedia Commons | CC BY-SA 4.0. Fotografía real de exhibición. Recorte para composición; conservar atribución y licencia. https://creativecommons.org/licenses/by-sa/4.0/ |
| fotografias/campo-rgb.png | Pix4D, R07 | Vista oficial de ortomosaico. Courtesy of Pix4D / pix4d.com. No calculada por este proyecto. |
| capturas/matic100.jpg | Pix4D, R07 | Vista oficial pública de nube y cámaras. Courtesy of Pix4D / pix4d.com. No acredita interfaz ni ejecución local. |
| capturas/fields-msp.png | Pix4D, R07 | Ejemplo público Rostock. Courtesy of Pix4D / pix4d.com. No corresponde al Field trial del manual. |

Fuente de las tres imágenes Pix4D: https://support.pix4d.com/hc/en-us/articles/360000235126. Enlace de atribución: https://www.pix4d.com/. Uso docente con atribución conforme a las condiciones de ejemplos. Los derechos permanecen en sus titulares. No se redistribuyen conjuntos ni cuadernos.
'''
(ROOT/'assets/CREDITOS.md').write_text(assets)
data={'matic100':{'url':'https://data.pix4d.com/misc/example_datasets/Pix4Dmatic_example_100_images.zip','crs':'EPSG:32632','alturas':'WGS84 elipsoidales','estado':'No descargado ni procesado'},'field_trial':{'url':'https://drive.google.com/file/d/1nyZhoSGoS7R0M_6q8sUcPblGiVXDZKU4/view?usp=drive_link','fuente':'M01 p.319','estado':'Acceso y apertura pendientes'},'targeted_operations':{'url':'https://drive.google.com/file/d/1uYEGeubg3l1e9mPq6w9Juk-AjHBj-s8s/view?usp=drive_link','fuente':'M01 p.301','estado':'Acceso y apertura pendientes'}}
for folder in ['ejercicios/pix4dmatic-pro','ejercicios/pix4dfields']:
    (ROOT/folder/'datos.json').write_text(json.dumps(data,ensure_ascii=False,indent=2))
    (ROOT/folder/'registro-ejecucion.csv').write_text('fecha,instructor,version,licencia,dataset,hash,crs,operacion,inicio,fin,resultado,evidencia,estado\n,,,,,,,,,,,,PENDIENTE\n')
    (ROOT/folder/'capturas-requeridas.md').write_text('# Capturas que debe aportar el ensayo\n\nRegistrar versión y proyecto; entradas, CRS y configuración; resultado e informe; comparación y exportación. Guardar PNG con fecha y origen. No se incluyen capturas inventadas ni páginas extensas del manual.\n')
(ROOT/'propuesta-curso/presupuesto.csv').write_text('rubro,cantidad,CRC_unitario,CRC_total\nDocencia,40,25000,1000000\nPreparacion,16,20000,320000\nAsistencia,24,12000,288000\nEquipo_y_operador,2,150000,300000\nLaboratorio,40,5000,200000\nMateriales,20,5000,100000\nTransporte,2,60000,120000\nSubtotal,,,2328000\nContingencia_10_por_ciento,,,232800\nBase_sin_costos_pendientes,,,2560800\n')
(ROOT/'LICENSE').write_text('Material original: Copyright 2026. Todos los derechos reservados hasta que el titular defina condiciones de distribución. El repositorio facilita consulta y uso docente del paquete; no concede derechos sobre marcas, software, manuales ni datos de terceros. Cada recurso conserva su licencia y atribución en assets/CREDITOS.md. No se presenta este archivo como licencia institucional UTN.\n')
(ROOT/'CHANGELOG.md').write_text('# Cambios\n\n## 2026-10-18\nReconstrucción y entrega de materiales desde fuente recuperada. Se distinguen archivos documentales y demostraciones operativas pendientes.\n\n## 2026-10-17\nInvestigación documental, revisión de manuales y selección de casos. La publicación de los binarios no se completó.\n')
(ROOT/'presentacion/fuentes/contenido.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2))
(ROOT/'taller/programa/programa.md').write_text('# Programa definitivo de trabajo · 120 minutos\n\n'+'\n'.join(f'- {x[0]} min: {x[2]}. Diapositivas {x[1]}.' for x in SECCIONES)+'\n\nFecha, hora, sede y expositor pendientes.\n')

if __name__=='__main__':print('Documentos generados: 7 DOCX; fuentes y registros preparados.')

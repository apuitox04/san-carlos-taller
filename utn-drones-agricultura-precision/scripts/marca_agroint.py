"""Aplica la identidad solicitada sin modificar el contenido técnico."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil,json
ROOT=Path(__file__).resolve().parents[1]
LOGO=ROOT/'assets/marca/logo_color_02.png'
for p in ROOT.rglob('*.docx'):
    if 'docs' in p.parts:continue
    d=Document(p)
    for style in d.styles:
        if hasattr(style,'font'):style.font.name='Fira Sans'
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        d.styles[name].font.color.rgb=RGBColor.from_string('264C6E')
    d.styles['Normal'].font.size=Pt(11)
    for paragraph in d.paragraphs:
        for run in paragraph.runs:
            if run.font.name:run.font.name='Fira Sans'
    for table in d.tables:
        for i,row in enumerate(table.rows):
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr()
                for shd in list(pr.findall(qn('w:shd'))):pr.remove(shd)
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'264C6E' if i==0 else ('F5F5F5' if i%2 else 'FFFFFF'));pr.append(shd)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name='Fira Sans';run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '263B48')
    for section in d.sections:
        section.header_distance=Cm(1.5);section.top_margin=Cm(4.4)
        head=section.header.paragraphs[0];head.clear();head.add_run().add_picture(str(LOGO),width=Cm(4))
        footer=section.footer.paragraphs[0];footer.clear()
        r=footer.add_run('AGROINT · Taller para Administración Agroindustrial UTN · agroint.co.cr');r.font.name='Fira Sans';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('264C6E')
    d.save(p)

doc=ROOT/'scripts/documentos.py'
s=doc.read_text().replace("NAVY='102F46'; TEAL='087F83'","NAVY='264C6E'; TEAL='00918C'").replace("'DejaVu Sans'","'Fira Sans'").replace("'102F46' if name!='Normal'","'264C6E' if name!='Normal'")
doc.write_text(s)

site=ROOT/'scripts/sitio.py';s=site.read_text()
s=s.replace('UTN <span>Administración Agroindustrial</span>','<img src="media/logo_color_02.png" alt="AGROINT"><span>Taller para la UTN</span>')
s=s.replace('<title>Drones y agricultura de precisión · UTN</title>','<title>Drones y agricultura de precisión · AGROINT para UTN</title>')
s=s.replace('De las fotografías a los mapas. De los mapas a una decisión que pueda comprobarse.','Productos geoespaciales para interpretar cultivos y fundamentar decisiones de gestión.')
s=s.replace('<footer><p>Universidad Técnica Nacional · Administración Agroindustrial</p>','<footer><img class="footerlogo" src="media/logo_color_02_blanco.png" alt="AGROINT"><p>Taller para la Universidad Técnica Nacional · Administración Agroindustrial</p><p><a href="https://agroint.co.cr/">agroint.co.cr</a></p>')
s=s.replace('#102f46','#264c6e').replace('#087f83','#00918c').replace('#6ce0d2','#dab055').replace('#eaf3f3','#f5f5f5').replace('system-ui,sans-serif','"Fira Sans",system-ui,sans-serif')
site.write_text(s)

brief='''# Aplicación de identidad AGROINT

Fecha: 18 octubre 2026. Solicitud del titular: presentar el taller con el estilo AGROINT.

## Fuentes entregadas
- Carpeta de marca: https://drive.google.com/drive/folders/1XWWStPmMFfeQ6MjkoGIxmnl3hFp7_TjR
- Libro de marca 2024: https://drive.google.com/file/d/12QUuRDSPeFFUOIp8h3lj03dg9lmfEhu2/view
- Logotipo color original: https://drive.google.com/file/d/1QNs3W_LPpZMtvMCkvh-2SLH54A59C41-/view
- Logotipo blanco original: https://drive.google.com/file/d/1CWufKKCLR1qx4wYgAbB36le8vah7N6IG/view
- Hoja membretada consultada: https://docs.google.com/document/d/19YAJE_dxAYWjPsPSYd7dbobkLSsgUei4/edit

Consulta y descarga: 18 octubre 2026. Se preservan los PNG originales, sin recoloración ni deformación. Los logotipos y la marca siguen siendo propiedad de su titular. Uso autorizado por la solicitud, sin concesión de derechos a terceros.

## Decisiones aplicadas
Fira Sans en textos y títulos. Se incluyen Regular y Bold con licencia OFL. Fuente: https://github.com/google/fonts/tree/main/ofl/firasans.

Paleta HEX del manual: #00918c, #264c6e, #7abb81 y #dab055. Neutros #F5F5F5 y blanco. El RGB impreso para #00918c no coincide con ese HEX; se adopta el HEX repetido en la sección web. No se altera el color incorporado en los logotipos originales.

Logotipo blanco sobre azul sólido y color sobre blanco. Portada y láminas internas: margen y protección de 60 px. En documentos se sigue la aplicación específica para membretes de 3–4 cm de ancho, con protección de 1.5 cm; el manual ofrece también un mínimo general de 7 cm para el horizontal, por lo que ambos criterios no son equivalentes. Se prioriza la aplicación específica para documentos.

La UTN permanece identificada como institución destinataria. No se incorpora un logotipo UTN ni se afirma una alianza, aprobación o patrocinio. Se conserva el contenido técnico y el estado pendiente de los ensayos Pix4D.
'''
(ROOT/'assets/marca/APLICACION.md').write_text(brief)
credit=ROOT/'assets/CREDITOS.md';credit.write_text(credit.read_text()+'\n## Identidad AGROINT\n\nLogotipos originales aportados por el usuario y utilizados a solicitud suya. Fuente, paleta y condiciones en marca/APLICACION.md. Fira Sans conserva la licencia OFL en fuentes/OFL.txt.\n')
readme=ROOT/'README.md';readme.write_text(readme.read_text()+'\n## Identidad AGROINT\n\nLa versión del 18 octubre 2026 aplica el libro de marca y los logotipos originales suministrados. Para editar con la tipografía prevista, instalar los TTF de assets/fuentes. Los PDF conservan la tipografía exportada. Tras regenerar documentos, ejecutar scripts/marca_agroint.py y repetir exportación y revisión. La validación operativa de Pix4D sigue pendiente.\n')
changes=ROOT/'CHANGELOG.md';changes.write_text(changes.read_text()+'\n## 2026-10-18 · Identidad AGROINT\n\nPaleta corporativa, Fira Sans y logotipos originales en presentación, afiches, documentos y sitio. Contenido técnico conservado. Libro de marca y membrete consultados.\n')
print('Identidad aplicada a documentos y fuentes del sitio.')
